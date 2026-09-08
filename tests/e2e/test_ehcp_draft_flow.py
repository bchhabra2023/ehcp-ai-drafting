import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
import urllib.request
from pathlib import Path

from docx import Document
from playwright.sync_api import sync_playwright


REPO_ROOT = Path(__file__).resolve().parents[2]
BACKEND_DIR = REPO_ROOT / "backend"
FRONTEND_DIR = REPO_ROOT / "frontend"
SIMPLE_INPUT_DIR = REPO_ROOT / "Test Cases" / "Simple Case Inputs"
BACKEND_OUTPUT_DIR = BACKEND_DIR / "output"
BACKEND_TEMP_DIR = BACKEND_DIR / "temp"
OUTPUT_FILENAME = "Ari Solven-Trail Draft EHCP.docx"


def wait_for_url(url: str, timeout_seconds: int = 120) -> None:
    deadline = time.time() + timeout_seconds
    last_error = None
    while time.time() < deadline:
        try:
            with urllib.request.urlopen(url, timeout=5) as response:
                if response.status == 200:
                    return
        except Exception as exc:  # pragma: no cover - diagnostic path
            last_error = exc
        time.sleep(1)
    raise AssertionError(f"Timed out waiting for {url}: {last_error}")


def kill_process_tree(process: subprocess.Popen) -> None:
    if process.poll() is not None:
        return
    subprocess.run(
        ["taskkill", "/PID", str(process.pid), "/T", "/F"],
        check=False,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )


def collect_output_docx() -> Path:
    matches = sorted(BACKEND_OUTPUT_DIR.rglob(OUTPUT_FILENAME))
    if len(matches) != 1:
        raise AssertionError(f"Expected exactly one generated output file, found: {matches}")
    return matches[0]


class EHCPDraftFlowE2ETest(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        shutil.rmtree(BACKEND_OUTPUT_DIR, ignore_errors=True)
        shutil.rmtree(BACKEND_TEMP_DIR, ignore_errors=True)

        cls.backend_log = tempfile.NamedTemporaryFile(
            prefix="ehcp-backend-", suffix=".log", delete=False
        )
        cls.frontend_log = tempfile.NamedTemporaryFile(
            prefix="ehcp-frontend-", suffix=".log", delete=False
        )

        backend_env = os.environ.copy()
        backend_env["AUTH_ENABLED"] = "false"
        backend_env["AUDIT_LOG_ENABLED"] = "false"
        backend_env["E2E_FIXTURE_MODE"] = "true"
        backend_env["E2E_FIXTURE_CASE"] = "simple"
        cls.backend_process = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "uvicorn",
                "main:app",
                "--host",
                "127.0.0.1",
                "--port",
                "8000",
            ],
            cwd=BACKEND_DIR,
            env=backend_env,
            stdout=cls.backend_log,
            stderr=subprocess.STDOUT,
        )

        frontend_env = os.environ.copy()
        frontend_env["AUTH_ENABLED"] = "false"
        frontend_env["BACKEND_URL"] = "http://127.0.0.1:8000"
        cls.frontend_process = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "streamlit",
                "run",
                "app.py",
                "--server.headless",
                "true",
                "--server.port",
                "8501",
                "--browser.gatherUsageStats",
                "false",
            ],
            cwd=FRONTEND_DIR,
            env=frontend_env,
            stdout=cls.frontend_log,
            stderr=subprocess.STDOUT,
        )

        wait_for_url("http://127.0.0.1:8000/api/health")
        wait_for_url("http://127.0.0.1:8501")

    @classmethod
    def tearDownClass(cls) -> None:
        kill_process_tree(cls.frontend_process)
        kill_process_tree(cls.backend_process)
        cls.backend_log.close()
        cls.frontend_log.close()

    def test_generates_draft_ehcp_from_simple_case_inputs(self) -> None:
        input_files = [
            SIMPLE_INPUT_DIR / "Personal Details - Ari Solven-Trail.docx",
            SIMPLE_INPUT_DIR / "Education Advice - A Trail.docx",
            SIMPLE_INPUT_DIR / "Health Advice - Ari Solven-Trail.pdf",
            SIMPLE_INPUT_DIR / "Social Care Advice - Ari S Trail.docx",
        ]

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(headless=True)
            page = browser.new_page()
            page.goto("http://127.0.0.1:8501", wait_until="domcontentloaded")
            page.get_by_text("Upload EHCP Documents").wait_for(timeout=120000)

            page.locator("input[type='file']").set_input_files([str(path) for path in input_files])
            page.get_by_role("button", name="Read and Analyse").wait_for(timeout=120000)
            proceed_button = page.get_by_role("button", name="Yes, proceed with all files")
            if proceed_button.count():
                proceed_button.click()
                page.get_by_role("button", name="Read and Analyse").wait_for(timeout=120000)
            page.get_by_role("button", name="Read and Analyse").click()

            page.get_by_role("button", name="Create Draft EHCP").wait_for(timeout=120000)
            page.get_by_role("button", name="Create Draft EHCP").click()

            page.get_by_text("Final EHCP Document").wait_for(timeout=120000)
            page.get_by_role("button", name="Download Draft EHCP").wait_for(timeout=120000)
            browser.close()

        output_docx = collect_output_docx()
        doc = Document(output_docx)
        text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_parts.append(cell.text)
        body_text = "\n".join(text_parts)

        self.assertIn("Ari Solven-Trail", body_text)
        self.assertIn("Riverbank Primary School", body_text)


if __name__ == "__main__":
    unittest.main()
